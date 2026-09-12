"""CSV and Excel export services for monthly aggregated loan records."""

import csv
import io
from typing import List
from src.db.models import LoanRecord

class CSVExporter:
    """Exports monthly loan records to RFC 4180 CSV and Excel formats."""

    HEADER = [
        "id",
        "serial_number",
        "borrower_name",
        "mobile_number",
        "address",
        "loan_amount",
        "ocr_engine_used",
        "verified",
        "verified_at",
        "created_at",
    ]

    def export_monthly_csv(self, loans: List[LoanRecord]) -> str:
        """
        Generate RFC 4180 compliant CSV string for a list of loan records.
        Appends summary total line if records are present.
        """
        output = io.StringIO()
        writer = csv.writer(
            output,
            dialect="excel",
            quoting=csv.QUOTE_MINIMAL,
            lineterminator="\r\n",  # Strict RFC 4180 CRLF line terminators
        )

        # Write header
        writer.writerow(self.HEADER)

        # Write data rows
        total_amount = 0.0
        for loan in loans:
            total_amount += float(loan.amount)
            writer.writerow(
                [
                    loan.id,
                    loan.serial_number,
                    loan.name,
                    loan.mobile,
                    loan.address or "",
                    f"{loan.amount:.2f}",
                    loan.ocr_engine_used,
                    "True" if loan.verified else "False",
                    loan.verified_at or "",
                    loan.created_at,
                ]
            )

        # If data rows exist, append summary total row
        if loans:
            writer.writerow(
                [
                    "TOTAL",
                    f"{len(loans)} records",
                    "",
                    "",
                    "TOTAL_AMOUNT",
                    f"{total_amount:.2f}",
                    "",
                    "",
                    "",
                    "",
                ]
            )

        return output.getvalue()

    def export_monthly_excel(self, loans: List[LoanRecord]) -> bytes:
        """
        Generate Excel (.xlsx) binary bytes for a list of loan records.
        Falls back to CSV bytes if openpyxl is unavailable.
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Loan Records"

            # Header style
            header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

            ws.append(self.HEADER)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            total_amount = 0.0
            for loan in loans:
                total_amount += float(loan.amount)
                ws.append(
                    [
                        loan.id,
                        loan.serial_number,
                        loan.name,
                        loan.mobile,
                        loan.address or "",
                        float(loan.amount),
                        loan.ocr_engine_used,
                        "True" if loan.verified else "False",
                        loan.verified_at or "",
                        loan.created_at,
                    ]
                )

            # Total row
            if loans:
                total_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
                total_font = Font(name="Arial", size=11, bold=True)
                total_row = ["TOTAL", f"{len(loans)} records", "", "", "TOTAL_AMOUNT", total_amount, "", "", "", ""]
                ws.append(total_row)
                for cell in ws[ws.max_row]:
                    cell.fill = total_fill
                    cell.font = total_font

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            buffer = io.BytesIO()
            wb.save(buffer)
            return buffer.getvalue()
        except Exception:
            # Fallback to UTF-8 encoded CSV bytes
            return self.export_monthly_csv(loans).encode("utf-8")
